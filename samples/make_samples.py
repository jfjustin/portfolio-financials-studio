"""Generate realistic sample reports to exercise every path of the pipeline.

Creates (in samples/):
  - fund_financials_q4.xlsx        (balance sheet + income statement that reconcile)
  - jv_holdings.xlsx               (joint-venture financials)
  - direct_investment.xlsx         (cost / fair value / MOIC that reconcile)
  - jv_equity_income.pdf           (NATIVE pdf; equity income = ownership × JV NI)

Run:  python samples/make_samples.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent


def _kv_sheet(path: Path, title: str, header: list[str], rows: list[tuple]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for line in header:
        ws.append([line])
    ws.append([])
    ws.append(["Line item", "Amount"])
    for label, value in rows:
        ws.append([label, value])
    wb.save(path)
    print("wrote", path.name)


def fund_financials():
    _kv_sheet(
        OUT / "fund_financials_q4.xlsx",
        "fund",
        ["Evergreen Capital Fund III, L.P.",
         "Statement of Assets, Liabilities and Net Assets",
         "As of December 31, 2024 (USD)"],
        [("Total Assets", 50_000_000),
         ("Total Liabilities", 20_000_000),
         ("Net Asset Value", 30_000_000),
         ("Total Investment Income", 4_000_000),
         ("Total Expenses", 1_500_000),
         ("Net Income", 2_500_000),
         ("Capital Called", 45_000_000),
         ("Distributions", 3_000_000),
         ("Management Fees", 600_000)],
    )


def jv_holdings():
    _kv_sheet(
        OUT / "jv_holdings.xlsx",
        "jv",
        ["Harbor Logistics Joint Venture",
         "JV Financial Summary — Year ended 31 December 2024 (USD)",
         "Our equity interest / ownership"],
        [("Ownership %", 40),
         ("Total Assets", 80_000_000),
         ("Total Liabilities", 30_000_000),
         ("Total Equity", 50_000_000),
         ("Revenue", 65_000_000),
         ("Net Income", 12_500_000),
         ("Carrying Value of Investment", 20_000_000)],
    )


def direct_investment():
    _kv_sheet(
        OUT / "direct_investment.xlsx",
        "direct",
        ["Direct Investment — Nimbus Technologies Inc.",
         "Portfolio holding valuation as of Q4 2024 (USD)"],
        [("Cost", 10_000_000),
         ("Fair Value", 18_000_000),
         ("Unrealized Gain", 8_000_000),
         ("Ownership %", 15),
         ("MOIC", 1.8),
         ("IRR", 22)],
    )


def jv_equity_income_pdf():
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 14)
    pdf.cell(0, 10, "Meridian Renewables Joint Venture", ln=True)
    pdf.set_font("helvetica", "", 11)
    pdf.cell(0, 7, "Equity Method Investment - Equity in Earnings", ln=True)
    pdf.cell(0, 7, "For the year ended December 31, 2024 (USD)", ln=True)
    pdf.ln(4)

    def leader(label, value):
        dots = "." * max(3, 60 - len(label))
        pdf.cell(0, 7, f"{label} {dots} {value}", ln=True)

    pdf.set_font("helvetica", "", 11)
    leader("Our ownership interest", "40%")
    leader("Net income for the year (100%)", "12,500,000")
    leader("Equity in earnings of joint venture", "5,000,000")
    leader("Dividends received", "1,200,000")
    leader("Investment carrying value", "20,000,000")
    out = OUT / "jv_equity_income.pdf"
    pdf.output(str(out))
    print("wrote", out.name)


if __name__ == "__main__":
    fund_financials()
    jv_holdings()
    direct_investment()
    jv_equity_income_pdf()
    print("\nSample files ready in", OUT)
